#!/usr/bin/env python
"""Build the metadata-driven TOPOPLACER teaser and editable PowerPoint.

使用示例:
    python tools/build_paper_teaser.py
"""

from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import subprocess
import sys

import numpy as np
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("MPLCONFIGDIR", str(ROOT / "outputs" / ".matplotlib"))

from tools import export_method_overview_materials as materials
from src.training.lc_bgplacenet_stage2 import build_space_former_targets

OUTPUT = ROOT / "outputs" / "paper_teaser"
SAMPLES = (
    "hope__scene_0000__0325__obj_3",
    "hope__scene_0006__0000__obj_4",
)
COLORS = {
    "ink": "202020", "muted": "555555", "source": "E69F00",
    "reference": "0072B2", "valid": "009E73", "invalid": "D55E00",
}


def write_json(path: Path, payload: dict) -> None:
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def save_support_map(assets, key, box, component, footprint, supported):
    """Show the evaluated cells in canonical XY, with the full footprint visible."""
    size = (640, 440)
    center = np.asarray(box[:2])
    scale = 29.0
    xy_to_uv = lambda xy: (np.asarray(xy) - center) * np.array([scale, -scale]) + np.array(size) / 2
    image = Image.new("RGB", size, "#F4F4F4")
    draw = ImageDraw.Draw(image)
    cells = np.vstack([component[:, :2], footprint[:, :2]])
    colors = ["#DCEBE5"] * len(component) + ["#A3CDBA" if ok else "#D55E00" for ok in supported]
    for uv, color in zip(xy_to_uv(cells), colors):
        x, y = uv
        if -scale < x < size[0] + scale and -scale < y < size[1] + scale:
            draw.rectangle((x - scale / 2 + 1, y - scale / 2 + 1, x + scale / 2 - 1, y + scale / 2 - 1), fill=color)
    corners = materials.place_box_to_corners(box)
    bottom = corners[np.isclose(corners[:, 2], corners[:, 2].min())]
    order = np.argsort(np.arctan2(bottom[:, 1] - center[1], bottom[:, 0] - center[0]))
    uv = xy_to_uv(bottom[order, :2])
    base = assets / f"{key}_base.png"
    image.save(base)
    draw.line([tuple(p) for p in np.vstack([uv, uv[:1]])], fill="#202020", width=4)
    cx, cy = np.array(size) / 2
    draw.line((cx - 8, cy, cx + 8, cy), fill="#202020", width=3)
    draw.line((cx, cy - 8, cx, cy + 8), fill="#202020", width=3)
    image.save(assets / f"{key}.png")
    return {"base": f"assets/{key}_base.png", "size": size,
            "boxes": [{"name": "Evaluated footprint", "uv": uv.tolist(), "color": "202020"}],
            "edges": [(0, 1), (1, 2), (2, 3), (3, 0)], "view": "canonical XY; +Y upward; cell size 1 cm"}


def render_assets(output: Path = OUTPUT) -> dict:
    """Render scientific overlays; preserve projected boxes for native PPT lines."""
    assets = output / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    cfg = materials.load_config(ROOT / "configs/lc_bgplacenet_stage2_enriched.yaml")
    cfg["data"]["sources"] = [s for s in cfg["data"]["sources"] if s["name"] == "hope"]
    items = materials.build_stage2_index(materials.build_sources_from_config(cfg))
    indices = [ROOT / "outputs" / "method_overview_materials" / name / "materials_index.json" for name in SAMPLES]
    records = [json.loads(path.read_text()) for path in indices]
    selected = [next(item for item in items if item.item_id == record["item_id"]) for record in records]
    item = selected[0]
    scene, points, _ = materials.load_scene_and_points(item)
    boxes = materials.scene_box_arrays(scene, item, None)
    context = materials.load_collision_context(item)
    visual_context = materials.canonical_scene_collision_context(scene)
    gt = boxes["final_box"]
    bad_language = np.asarray(records[0]["boxes"]["semantic_wrong_box"])
    bad_support = np.asarray(records[0]["boxes"]["partial_support_probe_box"])
    bad_collision = np.asarray(records[0]["boxes"]["collision_probe_box"])
    manifest = {"title": "TOPOPLACER", "data_kind": "GT and controlled probes; no model predictions", "panels": {}, "samples": [], "checks": {}}

    for row, entry, index_path in zip(selected, records, indices):
        manifest["samples"].append({
            "sample_id": row.sample_id, "item_id": row.item_id,
            "instruction": row.instruction, "target_relation": row.target_relation,
            "canonical_json": str(row.dataset_dir / "samples" / f"{row.sample_id}.json"),
            "yaw_set_npz": str(row.yaw_set_npz),
            "direction_filtered_heatmap_ply": str(row.direction_filtered_heatmap_ply),
            "source_materials_index": str(index_path.relative_to(ROOT)),
            "object_id": row.object_id, "reference_object_id": row.reference_object_id,
        })

    # Use the exact benchmark support band and retain source-location scene voxels.
    for key, box in (("gt", gt), ("language_probe", bad_language), ("support_probe", bad_support), ("collision_probe", bad_collision)):
        manifest["checks"][key] = {
            "box_xyz_lwh_yaw_cm_rad": box.tolist(),
            "relation": materials.placement_relation(scene, box, boxes["reference_corners"]),
            "support_coverage": materials.support_topology(points, box, 1.0, 3.0, 1.0)[2],
            **materials.compute_collision_metrics(box, context),
        }
    checks = manifest["checks"]
    assert checks["gt"]["relation"] == item.target_relation
    assert checks["gt"]["support_coverage"] == 1.0 and not checks["gt"]["collision"]
    assert checks["language_probe"]["relation"] != item.target_relation
    assert checks["language_probe"]["support_coverage"] == 1.0 and not checks["language_probe"]["collision"]
    assert checks["support_probe"]["relation"] == item.target_relation
    assert 0 < checks["support_probe"]["support_coverage"] < 1 and not checks["support_probe"]["collision"]
    assert checks["collision_probe"]["collision"]

    def panel(key, scene, image, wire_boxes):
        base = assets / f"{key}_base.png"
        image.convert("RGB").save(base)
        preview = image.convert("RGB").resize((image.width * 2, image.height * 2), Image.Resampling.LANCZOS)
        draw = ImageDraw.Draw(preview)
        projected = []
        for name, corners, color in wire_boxes:
            uv, _ = materials.project_world_points(scene, corners)
            projected.append({"name": name, "uv": uv.tolist(), "color": color})
            # Preview is raster; these same coordinates become editable PPT connectors.
            for a, b in materials.BOX_EDGES:
                segment = [tuple(2 * uv[a]), tuple(2 * uv[b])]
                draw.line(segment, fill="#FFFFFF", width=8)
                draw.line(segment, fill=f"#{color}", width=5)
        preview.save(assets / f"{key}.png")
        manifest["panels"][key] = {"base": str(base.relative_to(output)), "size": image.size, "boxes": projected}

    source = ("Source object", boxes["source_corners"], COLORS["source"])
    reference = ("Reference object", boxes["reference_corners"], COLORS["reference"])
    valid = ("GT placement", materials.place_box_to_corners(gt), COLORS["valid"])
    rgb = Image.fromarray(scene.rgb).convert("RGBA")
    panel("hero", scene, rgb, [source, reference, valid])
    panel("language_wrong", scene, rgb, [reference, ("Wrong relation", materials.place_box_to_corners(bad_language), COLORS["invalid"])])
    panel("language_valid", scene, rgb, [reference, valid])

    for key, box, color in (("support_wrong", bad_support, COLORS["invalid"]), ("support_valid", gt, COLORS["valid"])):
        image = rgb.copy()
        component, footprint, _, _ = materials.support_topology(points, box, 1, 3, 1)
        supported = materials.footprint_supported_mask(footprint, component)
        visible_component = component[~materials.points_inside_scene_obbs(component, visual_context, margin_cm=0)]
        materials.draw_projected_points(image, scene, visible_component, (23, 179, 142), radius=2, alpha=100, max_points=900)
        materials.draw_projected_points(image, scene, footprint[supported], (0, 173, 131), radius=3, alpha=220)
        materials.draw_projected_points(image, scene, footprint[~supported], (238, 105, 75), radius=4, alpha=240)
        panel(key, scene, image, [("Placement footprint", materials.place_box_to_corners(box), color)])
        manifest["panels"][key + "_xy"] = save_support_map(assets, key + "_xy", box, component, footprint, supported)

    collision_image = rgb.copy()
    surfaces = np.vstack([materials.obb_surface_sample_points(obb) for obb in visual_context["object_obbs"]])
    overlap = surfaces[materials.points_inside_yaw_box(surfaces, bad_collision)]
    materials.draw_projected_points(collision_image, scene, overlap, (238, 74, 56), radius=3, alpha=210, max_points=500)
    panel("collision_wrong", scene, collision_image, [reference, ("Controlled collision", materials.place_box_to_corners(bad_collision), COLORS["invalid"])])
    panel("collision_valid", scene, rgb, [reference, valid])

    # Direction filtering shares the training target builder; do not plot all yaw-set rows.
    multi = selected[1]
    multi_scene, multi_points, _ = materials.load_scene_and_points(multi)
    heat_points, scores, _ = materials.load_material_heatmap(multi, None)
    targets = build_space_former_targets(multi.yaw_set_npz, heat_points[scores > 0], 1.0)
    centers = targets["gt_bottom_centers"]
    with np.load(multi.yaw_set_npz) as yaw_data:
        raw_centers = yaw_data["bottom_center_world"]
        raw_angles = yaw_data["yaw_angles_rad"]
        raw_mask = yaw_data["valid_yaw_mask"]
    # Exact nearest correspondence only transfers masks after canonical target filtering.
    rows = np.argmin(np.linalg.norm(centers[:, None, :] - raw_centers[None, :, :], axis=2), axis=1)
    assert np.allclose(centers, raw_centers[rows], atol=1e-4)
    candidates = []
    multi_context = materials.load_collision_context(multi)
    multi_boxes = materials.scene_box_arrays(multi_scene, multi, None)
    for idx in np.argsort(raw_mask[rows].sum(axis=1), kind="stable"):
        if candidates and (np.linalg.norm(centers[idx] - candidates[0]["center"]) < 8 or raw_mask[rows[idx]].sum() < 3):
            continue
        allowed_angles = raw_angles[raw_mask[rows[idx]]]
        if candidates:
            difference = np.abs((allowed_angles - candidates[0]["box"][6] + np.pi / 2) % np.pi - np.pi / 2)
            angle = allowed_angles[np.argmax(difference)]
        else:
            angle = allowed_angles[0]
        candidate = materials.candidate_from_bottom_center(centers[idx], multi.place_box_gt[3:6], angle)
        if not materials.project_world_points(multi_scene, materials.place_box_to_corners(candidate))[1].all():
            continue
        if materials.support_topology(multi_points, candidate, 1, 3, 1)[2] != 1.0:
            continue
        if not materials.is_collision_free(candidate, multi_context):
            continue
        if materials.placement_relation(multi_scene, candidate, multi_boxes["reference_corners"]) != multi.target_relation:
            continue
        candidates.append({"center": centers[idx].tolist(), "box": candidate.tolist(), "yaw_angles_rad": raw_angles[raw_mask[rows[idx]]].tolist()})
        if len(candidates) == 2:
            break
    assert len(candidates) == 2
    manifest["multi_target"] = {"center_count": len(centers), "centers_cm": centers.tolist(), "selected": candidates, "note": "Selected boxes additionally checked with current benchmark metrics; dense dots show direction-filtered GT annotations."}
    multi_image = Image.fromarray(multi_scene.rgb).convert("RGBA")
    materials.draw_projected_points(multi_image, multi_scene, centers, (0, 114, 178), radius=3, alpha=175)
    multi_wire = [("Source object", multi_boxes["source_corners"], COLORS["source"]),
                  ("Reference object", multi_boxes["reference_corners"], COLORS["reference"])]
    panel("multi_centers", multi_scene, multi_image, multi_wire)
    manifest["multi_target"]["selected_uv"] = materials.project_world_points(multi_scene, np.asarray([c["center"] for c in candidates]))[0].tolist()
    alternatives = [(f"Alternative {i + 1}", materials.place_box_to_corners(np.asarray(c["box"])), COLORS["valid"] if i == 0 else COLORS["reference"]) for i, c in enumerate(candidates)]
    panel("multi_boxes", multi_scene, Image.fromarray(multi_scene.rgb).convert("RGBA"), alternatives)
    manifest["support_definition"] = "1 cm voxels; bottom -3/+1 cm; 3x3 closing, hole filling, 8-connected center component; retain source voxels"
    write_json(output / "teaser_metadata.json", manifest)
    sheet = Image.new("RGB", (1440, 398 * ((len(manifest["panels"]) + 2) // 3)), "white")
    draw = ImageDraw.Draw(sheet)
    for i, key in enumerate(manifest["panels"]):
        x, y = (i % 3) * 480, (i // 3) * 398
        draw.text((x + 12, y + 8), key.replace("_", " "), font=materials.load_font(20, bold=True), fill="#20343E")
        with Image.open(assets / f"{key}.png") as image:
            image.thumbnail((464, 348), Image.Resampling.LANCZOS)
            sheet.paste(image, (x + 8, y + 39))
    sheet.save(output / "materials_preview.png")
    return manifest


def build_presentation(output: Path, manifest: dict) -> Path:
    """Compose a compact paper figure; geometry carries color, all framing is white."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(18), Inches(6.6)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    def remove_effects(shape):
        # The default theme adds shadows even to lines and tiny swatches.
        style = shape._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
        if style is not None:
            shape._element.remove(style)
        shape._element.spPr.append(OxmlElement("a:effectLst"))

    def text(x, y, w, h, content, size=18, color="ink", bold=False, align="left"):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, value in enumerate(content.split("\n")):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = value
            para.alignment = PP_ALIGN.CENTER if align == "center" else PP_ALIGN.LEFT
            para.font.name = "DejaVu Sans"
            para.font.size = Pt(size)
            para.font.bold = bold
            para.font.color.rgb = RGBColor.from_string(COLORS.get(color, color))
            para.space_after = Pt(0)
        return box

    def line(x1, y1, x2, y2, color="ink", width=1.1, arrow=False):
        shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        remove_effects(shape)
        shape.line.color.rgb = RGBColor.from_string(COLORS.get(color, color))
        shape.line.width = Pt(width)
        if arrow:
            end = OxmlElement("a:tailEnd")
            end.set("type", "triangle")
            shape.line._get_or_add_ln().append(end)
        return shape

    def circle(x, y, radius, color, fill=False):
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x - radius), Inches(y - radius), Inches(2 * radius), Inches(2 * radius))
        remove_effects(shape)
        shape.line.color.rgb = RGBColor.from_string(COLORS.get(color, color))
        shape.line.width = Pt(1)
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(COLORS.get(color, color))
        else:
            shape.fill.background()
        return shape

    def image_panel(key, x, y, w, h, crop=None):
        entry = manifest["panels"][key]
        iw, ih = entry["size"]
        left, top, right, bottom = crop or (0, 0, iw, ih)
        picture = slide.shapes.add_picture(str(output / entry["base"]), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
        picture.crop_left, picture.crop_top = left / iw, top / ih
        picture.crop_right, picture.crop_bottom = 1 - right / iw, 1 - bottom / ih
        for wire in entry["boxes"]:
            uv = np.asarray(wire["uv"]) - np.array([left, top])
            for a, b in entry.get("edges", materials.BOX_EDGES):
                segment = clip_segment(uv[a], uv[b], right - left, bottom - top)
                if segment is None:
                    continue
                start, end = segment
                coords = (x + start[0] / (right - left) * w, y + start[1] / (bottom - top) * h,
                          x + end[0] / (right - left) * w, y + end[1] / (bottom - top) * h)
                line(*coords, "FFFFFF", 2.5)
                edge = line(*coords, wire["color"], 1.8)
                edge.name = f"{key}: {wire['name']}"
        if key.endswith("_xy"):
            line(x + w / 2 - .065, y + h / 2, x + w / 2 + .065, y + h / 2)
            line(x + w / 2, y + h / 2 - .065, x + w / 2, y + h / 2 + .065)

    # (a) One large scene anchors the task, with no decorative container.
    text(.18, .10, 5.60, .75, "Put the tomato sauce can\nback-left of the cream cheese box.", 20)
    image_panel("hero", .18, 1.04, 5.60, 4.20)
    for x, color, label in ((.25, "source", "Source"), (1.92, "reference", "Reference"), (4.01, "valid", "Placement")):
        line(x, 5.56, x + .22, 5.56, color, 2.5)
        text(x + .30, 5.38, 1.50, .35, label, 16)
    text(.18, 5.98, 5.60, .38, "(a) Instruction-guided re-placement", 19, bold=True, align="center")

    # (b) Show full support footprints and tightly cropped collision geometry.
    text(6.10, .13, 5.62, .40, "Topology-aware box reasoning", 21, bold=True, align="center")
    text(6.10, .61, 2.72, .30, "Support (top view)", 17, align="center")
    text(9.00, .61, 2.72, .30, "Collision (detail)", 17, align="center")
    image_panel("support_wrong_xy", 6.10, 1.04, 2.72, 1.87)
    image_panel("collision_wrong", 9.00, 1.04, 2.72, 1.87, (40, 85, 440, 360))
    coverage = manifest["checks"]["support_probe"]["support_coverage"]
    text(6.10, 2.96, 2.72, .32, f"Partial support: {coverage:.0%}", 17, "invalid", align="center")
    text(9.00, 2.96, 2.72, .32, "Object intersection", 17, "invalid", align="center")
    line(7.46, 3.31, 7.46, 3.56, "muted", 1.4, arrow=True)
    line(10.36, 3.31, 10.36, 3.56, "muted", 1.4, arrow=True)
    image_panel("support_valid_xy", 6.10, 3.64, 2.72, 1.87)
    image_panel("collision_valid", 9.00, 3.64, 2.72, 1.87, (40, 85, 440, 360))
    text(6.10, 5.57, 2.72, .32, "Full support: 100%", 17, "valid", align="center")
    text(9.00, 5.57, 2.72, .32, "Collision-free", 17, "valid", align="center")
    text(6.10, 5.98, 5.62, .38, "(b) Physically feasible placement", 19, bold=True, align="center")

    # (c) The arrows show the actual yaw sets of the two highlighted centers.
    text(12.11, .10, 5.65, .75, "Put the butter box\nfront-right of the cherries can.", 20)
    image_panel("multi_centers", 12.36, 1.04, 5.20, 3.90)
    for i, (uv, selected) in enumerate(zip(manifest["multi_target"]["selected_uv"], manifest["multi_target"]["selected"])):
        name = chr(65 + i)
        x, y = 12.36 + uv[0] / 640 * 5.20, 1.04 + uv[1] / 480 * 3.90
        circle(x, y, .065, "ink", fill=True)
        text(x - .36, y - .19, .28, .31, name, 17, "ink", True)
        gx, gy = 12.70 + i * 2.58, 5.60
        circle(gx, gy, .23, "AAAAAA")
        line(gx - .26, gy, gx + .26, gy, "CCCCCC", .7)
        line(gx, gy - .26, gx, gy + .26, "CCCCCC", .7)
        for angle in selected["yaw_angles_rad"]:
            line(gx, gy, gx + .23 * np.cos(angle), gy - .23 * np.sin(angle), "reference", 2, arrow=True)
        count = len(selected["yaw_angles_rad"])
        text(gx + .43, gy - .21, 1.66, .42, f"{name}: {count} yaw" + ("s" if count > 1 else ""), 18)
    count = manifest["multi_target"]["center_count"]
    text(12.11, 4.98, 5.65, .30, f"{count} direction-valid centers; center-specific yaw", 16, align="center")
    text(12.11, 5.98, 5.65, .38, "(c) Dense center–yaw annotations", 19, bold=True, align="center")
    text(6.10, 6.38, 11.66, .20, "GT illustrations; failed placements are controlled probes. Yaw diagrams use canonical XY axes.", 11, "muted", align="center")
    slide.notes_slide.notes_text_frame.text = (
        "Scientific teaser for TOPOPLACER. Observations and geometry derive from dataset metadata; "
        "no model predictions are shown. Failed placements are controlled probes. "
        "Support maps show benchmark connected components and footprint cells in canonical XY, "
        "with 1 cm cells and +Y upward. Collision panels use the same RGB crop. "
        "A and B refer to annotated centers; radial arrows show their raw yaw sets in canonical XY. "
        "See teaser_metadata.json for original instructions, sources and metric checks. "
        "Text, box edges, footprint outlines, labels, arrows and yaw diagrams are editable PowerPoint objects."
    )
    path = output / "topoplacer_teaser.pptx"
    prs.save(path)
    return path


def clip_segment(start, end, width, height):
    """Liang–Barsky clipping keeps native PowerPoint edges within RGB panels."""
    delta = end - start
    lower, upper = 0.0, 1.0
    for p, q in zip((-delta[0], delta[0], -delta[1], delta[1]), (start[0], width - start[0], start[1], height - start[1])):
        if abs(p) < 1e-12:
            if q < 0:
                return None
            continue
        ratio = q / p
        if p < 0:
            lower = max(lower, ratio)
        else:
            upper = min(upper, ratio)
    if lower > upper:
        return None
    return start + lower * delta, start + upper * delta


def export_print_formats(output: Path, pptx_path: Path) -> None:
    """Use LibreOffice to render the actual PPT, retaining vector text in PDF."""
    profile = (output / ".libreoffice_profile").resolve().as_uri()
    subprocess.run(["libreoffice", f"-env:UserInstallation={profile}", "--headless", "--convert-to", "pdf", "--outdir", str(output), str(pptx_path)], check=True, timeout=90)
    pdf = output / "topoplacer_teaser.pdf"
    if not pdf.exists():
        raise RuntimeError("LibreOffice did not export the PDF")
    subprocess.run(["pdftoppm", "-png", "-r", "200", "-singlefile", str(pdf), str(output / "topoplacer_teaser")], check=True, timeout=90)
    with Image.open(output / "topoplacer_teaser.png") as image:
        image.resize((1800, 660), Image.Resampling.LANCZOS).save(output / "preview.png")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=OUTPUT)
    args = parser.parse_args()
    os.chdir(ROOT)
    manifest = render_assets(args.output_dir)
    pptx_path = build_presentation(args.output_dir, manifest)
    export_print_formats(args.output_dir, pptx_path)
    print(pptx_path)


if __name__ == "__main__":
    main()
